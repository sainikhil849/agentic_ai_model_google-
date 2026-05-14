"""
District Multi-City Real Data Scraper - VERIFIED WORKING CODE
Uses proven DistrictAgent logic that worked for Hyderabad
Adapts it for multiple cities: Mumbai, Delhi, Chennai, Pune
"""

import time
import json
import re
import logging
import os
import sys
from typing import List, Dict, Optional
from datetime import datetime
from playwright.sync_api import sync_playwright

logging.basicConfig(
    level=logging.INFO,
    format='%(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

class DistrictMultiCityScraper:
    """Scrape District.in events for multiple cities"""
    
    # City slug mapping for District URLs
    CITY_SLUG_MAP = {
        "mumbai": "mumbai",
        "delhi": "delhi",
        "pune": "pune",
        "chennai": "chennai",
        "bengaluru": "bengaluru",
        "bangalore": "bengaluru",
        "hyderabad": "hyderabad",
    }
    
    def __init__(self):
        self.seen_urls = set()
    
    def get_listing_urls(self, city: str) -> List[str]:
        """Get priority URLs for a city on District"""
        city_slug = self.CITY_SLUG_MAP.get(city.lower(), city.lower())
        city_display = city.capitalize()
        
        # Priority order for finding events
        urls = [
            f"https://www.district.in/events/{city_slug}-ticket-booking",
            f"https://www.district.in/events/{city_slug}",
            f"https://www.district.in/activities/{city_slug}-activities",
            f"https://www.district.in/events/",
        ]
        return urls
    
    def deep_scroll(self, page, scrolls: int = 15):
        """Scroll page to load dynamic content"""
        for i in range(scrolls):
            page.evaluate("window.scrollBy(0, window.innerHeight)")
            time.sleep(0.3)
            if i % 5 == 0:
                logger.debug(f"  Scroll {i+1}/{scrolls}")
    
    def extract_json_ld(self, page) -> List[Dict]:
        """Extract JSON-LD structured data from page"""
        scripts = page.query_selector_all('script[type="application/ld+json"]')
        results = []
        for script in scripts:
            try:
                text = script.inner_text()
                if text:
                    results.append(json.loads(text))
            except Exception:
                continue
        return results
    
    def get_event_candidates(self, page, city: str, target_count: int) -> List[Dict]:
        """Collect event URL candidates from listing pages"""
        candidates = []
        seen = set()
        city_slug = self.CITY_SLUG_MAP.get(city.lower(), city.lower())
        city_display = city.capitalize()
        
        urls = self.get_listing_urls(city)
        
        for url_idx, url in enumerate(urls):
            if len(candidates) >= target_count * 3:
                break
            
            try:
                logger.info(f"\n📍 Scraping URL {url_idx + 1}/{len(urls)}: {url}")
                page.goto(url, wait_until="domcontentloaded", timeout=60000)
                time.sleep(1.0)
                
                # Scroll to load all events
                self.deep_scroll(page, scrolls=15)
                
                # Find all event links
                anchors = page.query_selector_all("a[href*='/events/'], a[href*='/activities/']")
                url_count = 0
                
                for anchor in anchors:
                    if len(candidates) >= target_count * 3:
                        break
                    
                    href = anchor.get_attribute("href") or ""
                    if not href:
                        continue
                    
                    # Build full URL
                    full_url = href if href.startswith("http") else f"https://www.district.in{href}"
                    full_url = full_url.split("?")[0].split("#")[0]
                    
                    # Skip if already seen
                    if full_url in seen:
                        continue
                    
                    # Skip invalid URLs
                    if full_url.endswith(("/events", "/events/", "/activities", "/activities/")):
                        continue
                    if "javascript:" in full_url or "mailto:" in full_url:
                        continue
                    if "/events/" not in full_url and "/activities/" not in full_url:
                        continue
                    
                    seen.add(full_url)
                    candidates.append({
                        "url": full_url,
                        "city": city_display,
                        "city_slug": city_slug,
                    })
                    url_count += 1
                
                logger.info(f"  ✓ Got {url_count} candidates from this URL")
                
            except Exception as e:
                logger.debug(f"  ✗ Error on URL: {e}")
                continue
        
        logger.info(f"\n✓ Total candidates collected: {len(candidates)}")
        return candidates
    
    def extract_event_details(self, page, candidate: Dict) -> Optional[Dict]:
        """Extract event details from event detail page"""
        try:
            url = candidate.get("url", "")
            city = candidate.get("city", "")
            
            # Navigate to event page
            resp = page.goto(url, wait_until="domcontentloaded", timeout=15000)
            
            if not resp or resp.status >= 400:
                logger.debug(f"  ✗ 404 or error: {resp.status if resp else 'unknown'}")
                return None
            
            time.sleep(0.5)
            
            # Extract JSON-LD data
            event_data = {
                "event_name": None,
                "event_date": None,
                "event_time": None,
                "price": None,
                "venue": None,
                "event_url": url,
                "description": None,
                "organizer": None,
                "location": city,
                "platform": "District",
            }
            
            # Parse JSON-LD
            for ld in self.extract_json_ld(page):
                ld_type = str(ld.get("@type", "")).lower()
                
                if "event" in ld_type:
                    # Event name
                    if ld.get("name"):
                        event_data["event_name"] = ld.get("name")
                    
                    # Date and Time
                    if ld.get("startDate"):
                        start_dt = ld.get("startDate", "")
                        if "T" in start_dt:
                            # ISO format: 2026-04-25T20:30:00
                            parts = start_dt.split("T")
                            event_data["event_date"] = parts[0]
                            if len(parts) > 1:
                                event_data["event_time"] = parts[1][:5]  # HH:MM
                        else:
                            event_data["event_date"] = start_dt
                    
                    # Description
                    if ld.get("description"):
                        event_data["description"] = ld.get("description")[:200]
                    
                    # Venue/Location
                    location_data = ld.get("location", {})
                    if isinstance(location_data, dict):
                        if location_data.get("name"):
                            event_data["venue"] = location_data.get("name")
                    
                    # Price
                    offers = ld.get("offers", {})
                    if isinstance(offers, list) and offers:
                        offers = offers[0]
                    if isinstance(offers, dict):
                        if offers.get("price"):
                            try:
                                event_data["price"] = int(float(str(offers.get("price")).replace(",", "")))
                            except:
                                pass
                    
                    # Organizer
                    if ld.get("organizer"):
                        org = ld.get("organizer")
                        if isinstance(org, dict):
                            event_data["organizer"] = org.get("name")
                        else:
                            event_data["organizer"] = str(org)
                    
                    break
            
            # Fallback: Extract from page text if JSON-LD didn't have all data
            body_text = page.inner_text("body") or ""
            
            # Extract price from body if not found in JSON-LD
            if not event_data.get("price") and body_text:
                price_match = re.search(r"₹\s*([0-9,]+)", body_text)
                if price_match:
                    try:
                        event_data["price"] = int(price_match.group(1).replace(",", ""))
                    except:
                        pass
            
            # Extract time from body if not found
            if not event_data.get("event_time") and body_text:
                time_match = re.search(r"(\d{1,2}):(\d{2})\s*(AM|PM|am|pm)?", body_text)
                if time_match:
                    event_data["event_time"] = f"{time_match.group(1)}:{time_match.group(2)}"
            
            # Validation - must have these fields
            if not event_data.get("event_name"):
                logger.debug(f"  ✗ Missing event name")
                return None
            
            if not event_data.get("event_date"):
                logger.debug(f"  ✗ Missing event date")
                return None
            
            logger.info(f"  ✓ EXTRACTED: {event_data['event_name'][:50]}")
            logger.info(f"      Date: {event_data['event_date']} | Price: ₹{event_data.get('price', '?')}")
            logger.info(f"      Venue: {event_data.get('venue', 'N/A')[:40]}")
            
            return event_data
            
        except Exception as e:
            logger.debug(f"  ✗ Extraction error: {e}")
            return None
    
    def scrape_city(self, city: str, target_count: int = 100) -> List[Dict]:
        """Scrape events for a single city"""
        logger.info("\n" + "="*80)
        logger.info(f"DISTRICT - SCRAPING: {city.upper()}")
        logger.info("="*80)
        
        events = []
        
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            context = browser.new_context(
                user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
            )
            page = context.new_page()
            
            try:
                # Step 1: Get candidates
                candidates = self.get_event_candidates(page, city, target_count)
                
                if not candidates:
                    logger.warning(f"  ✗ No candidates found for {city}")
                    return []
                
                # Step 2: Extract details from each candidate
                logger.info(f"\nExtracting details from {len(candidates)} candidates...")
                
                for idx, candidate in enumerate(candidates):
                    if len(events) >= target_count:
                        break
                    
                    try:
                        # Create new page for each event to avoid context issues
                        detail_page = context.new_page()
                        
                        logger.info(f"\n  Event {idx + 1}/{len(candidates)}")
                        event = self.extract_event_details(detail_page, candidate)
                        
                        if event:
                            events.append(event)
                        
                        detail_page.close()
                        time.sleep(0.3)
                        
                    except Exception as e:
                        logger.debug(f"    Error processing candidate: {e}")
                        try:
                            detail_page.close()
                        except:
                            pass
                        continue
                
                logger.info(f"\n✓ Successfully extracted {len(events)} events from {city}")
                
            finally:
                browser.close()
        
        return events
    
    def scrape_all_cities(self, cities: List[str] = None, target_per_city: int = 100):
        """Scrape events for all cities"""
        if not cities:
            cities = ["Mumbai", "Delhi", "Chennai", "Pune"]
        
        all_events = {}
        
        for city in cities:
            events = self.scrape_city(city, target_per_city)
            all_events[city] = events
        
        return all_events
    
    def save_to_json(self, all_events: Dict[str, List[Dict]]):
        """Save events to JSON"""
        os.makedirs("exports", exist_ok=True)
        
        # Save consolidated
        consolidated_file = "exports/district_multi_city_real.json"
        with open(consolidated_file, "w", encoding="utf-8") as f:
            json.dump(all_events, f, indent=2, ensure_ascii=False)
        logger.info(f"✓ Saved: {consolidated_file}")
        
        # Save individual cities
        for city, events in all_events.items():
            city_file = f"exports/district_{city.lower()}_real.json"
            with open(city_file, "w", encoding="utf-8") as f:
                json.dump({city: events}, f, indent=2, ensure_ascii=False)
            logger.info(f"✓ Saved: {city_file}")
    
    def export_to_excel(self, all_events: Dict[str, List[Dict]]):
        """Export to Excel"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            
            os.makedirs("exports", exist_ok=True)
            excel_file = "exports/district_multi_city_real.xlsx"
            
            wb = Workbook()
            wb.remove(wb.active)
            
            # Summary sheet
            ws_summary = wb.create_sheet("Summary", 0)
            ws_summary["A1"] = "District Multi-City Events - Real Data"
            ws_summary["A1"].font = Font(bold=True, size=14)
            
            row = 3
            for city, events in all_events.items():
                ws_summary[f"A{row}"] = city
                ws_summary[f"B{row}"] = len(events)
                row += 1
            
            # City sheets
            for city, events in all_events.items():
                ws = wb.create_sheet(city)
                
                headers = ["Event Name", "Date", "Time", "Price (₹)", "Venue", "Location", "Event URL", "Description", "Organizer"]
                
                for col_idx, header in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.value = header
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                
                for row_idx, event in enumerate(events, 2):
                    ws.cell(row=row_idx, column=1).value = event.get("event_name", "")
                    ws.cell(row=row_idx, column=2).value = event.get("event_date", "")
                    ws.cell(row=row_idx, column=3).value = event.get("event_time", "")
                    ws.cell(row=row_idx, column=4).value = event.get("price", "")
                    ws.cell(row=row_idx, column=5).value = event.get("venue", "")
                    ws.cell(row=row_idx, column=6).value = event.get("location", "")
                    ws.cell(row=row_idx, column=7).value = event.get("event_url", "")
                    ws.cell(row=row_idx, column=8).value = event.get("description", "")
                    ws.cell(row=row_idx, column=9).value = event.get("organizer", "")
                
                # Adjust column widths
                ws.column_dimensions["A"].width = 35
                ws.column_dimensions["B"].width = 12
                ws.column_dimensions["C"].width = 10
                ws.column_dimensions["D"].width = 12
                ws.column_dimensions["E"].width = 30
                ws.column_dimensions["F"].width = 15
                ws.column_dimensions["G"].width = 50
                ws.column_dimensions["H"].width = 35
                ws.column_dimensions["I"].width = 25
            
            wb.save(excel_file)
            logger.info(f"✓ Excel saved: {excel_file}")
            return excel_file
            
        except Exception as e:
            logger.error(f"Excel export error: {e}")
            return None

def main():
    logger.info("\n" + "="*80)
    logger.info("DISTRICT.IN MULTI-CITY REAL DATA SCRAPER")
    logger.info("Using proven District scraping logic")
    logger.info("="*80 + "\n")
    
    scraper = DistrictMultiCityScraper()
    
    # Scrape all cities
    all_events = scraper.scrape_all_cities(
        cities=["Mumbai", "Delhi", "Chennai", "Pune"],
        target_per_city=100
    )
    
    # Summary
    logger.info("\n" + "="*80)
    logger.info("SCRAPING COMPLETE - SUMMARY")
    logger.info("="*80)
    
    total = 0
    for city, events in all_events.items():
        count = len(events)
        total += count
        logger.info(f"{city}: {count} real events scraped from District.in")
    
    logger.info(f"\nTOTAL: {total} real events")
    
    # Save to JSON and Excel
    logger.info("\nSaving to files...")
    scraper.save_to_json(all_events)
    scraper.export_to_excel(all_events)
    
    logger.info("\n✓ COMPLETE - All files saved to exports/")
    logger.info("="*80 + "\n")

if __name__ == "__main__":
    main()
